import json
from pathlib import Path
import tempfile
import unittest

from analyze_matched_release import HOSTS, analyze, change, word_document
from build_arm_compilers import digest
from matched_release import RUST, LLVM, PERF, VARIANTS
from native_ci_probe import NATIVE_COMMANDS, paired_order


def fixture(root):
    releases, benchmarks = root / "releases", root / "benchmarks"
    releases.mkdir()
    benchmarks.mkdir()
    protocol = {"rust_sha": RUST, "llvm_sha": LLVM, "perf_sha": PERF, "source_sha": "workload",
                "cores": 4, "release_repetitions_per_cell": 3, "warmup_rounds": 1, "measured_rounds": 3}
    jobs = {"jobs": []}
    for arch, host in HOSTS.items():
        for vm in range(1, 4):
            directory = benchmarks / f"{arch}-{vm}"
            directory.mkdir()
            tools = {"msvc": "same", "sdk": "same", "files": {}}
            index = {"host": host, "vm": vm, "complete": True, "pilot": False, "protocol": protocol,
                     "release_run": "release", "benchmark_run": "benchmark", "benchmark_attempt": "1",
                     "tools": tools, "compiler_metadata": {}, "artifact_evidence": {}, "runs": []}
            for number, variant in enumerate(VARIANTS):
                folder = releases / f"release-{arch}-{variant}-{vm}-evidence"
                folder.mkdir(parents=True)
                metadata = {
                    "complete": True, "pilot": False, "host": host, "variant": variant, "vm": vm,
                    "attempt": "1", "rust_sha": RUST, "llvm_sha": LLVM, "perf_sha": PERF,
                    "run_id": "release", "tools": tools,
                }
                (folder / "result.json").write_text(json.dumps(metadata))
                (folder / "distribution-timing.json").write_text(json.dumps({"wall_seconds": 80}))
                (folder / "extracted-dist-tests.json").write_text(json.dumps({"wall_seconds": 10, "exit_code": 0}))
                archives = {f"rustc-nightly-{host}.tar.xz": "compiler"}
                if variant == "baseline":
                    archives.update({f"cargo-nightly-{host}.tar.xz": "cargo", f"rust-std-nightly-{host}.tar.xz": "std"})
                (folder / "dist-manifest.json").write_text(json.dumps(
                    [{"name": name, "sha256": checksum} for name, checksum in archives.items()]))
                index["compiler_metadata"][variant] = metadata
                index["artifact_evidence"][variant] = {
                    "result_sha256": digest(folder / "result.json"), "archives": archives}
                jobs["jobs"].append({
                    "id": len(jobs["jobs"]) + 1, "name": f"{arch} / {variant} / {vm}",
                    "run_id": "release", "run_attempt": 1, "status": "completed", "conclusion": "success",
                    "started_at": "2026-10-11T10:00:00Z",
                    "completed_at": ["2026-10-11T10:01:40Z", "2026-10-11T10:02:30Z",
                                     "2026-10-11T10:03:20Z"][number],
                })
            for r in range(4):
                for position, variant in enumerate(paired_order(list(VARIANTS), vm, r)):
                    path = f"{r}-{variant}.json"
                    block = {
                        "metadata": {"source_sha": "workload", "variant": variant, "round": r,
                                     "profile_environment": {"CARGO_NET_OFFLINE": "true", "CARGO_BUILD_JOBS": "4"}},
                        "measurements": [
                            {"name": name, "exit_code": 0, "wall_seconds": {"baseline": 10, "msvc-pgo": 8,
                                                                         "llvm-pgo-thin": 7}[variant],
                             "cpu_seconds": 20, "affinity": [0, 1, 2, 3], "compiled_artifacts": 1,
                             "command": ["cargo", *args, "--frozen"]}
                            for name, args in NATIVE_COMMANDS],
                        "correctness": [{"name": name, "exit_code": 0,
                                         "environment_overrides": {"RUST_TEST_THREADS": "1"}}
                                        for name, _ in NATIVE_COMMANDS[1:]] if r == 0 else [],
                    }
                    (directory / path).write_text(json.dumps(block))
                    index["runs"].append({"variant": variant, "round": r, "position": position,
                                          "warmup": r == 0, "path": path})
            (directory / "index.json").write_text(json.dumps(index))
    return releases, benchmarks, jobs, protocol


class MatchedAnalysisTests(unittest.TestCase):
    def test_rejects_selfhosted_release_in_hosted_comparison(self):
        with tempfile.TemporaryDirectory() as tmp:
            args = fixture(Path(tmp))
            path = args[0] / "release-arm64-baseline-1-evidence/result.json"
            metadata = json.loads(path.read_text())
            for runner_class in ("self-hosted-arm64", "self-hosted-x64", "self-hosted-arm64-highcore"):
                with self.subTest(runner_class=runner_class):
                    metadata["runner_class"] = runner_class
                    path.write_text(json.dumps(metadata))
                    with self.assertRaisesRegex(ValueError, "self-hosted results"):
                        analyze(*args)

    def test_statistics_and_complete_document(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            result = analyze(*fixture(root))
            for rows in result["architectures"].values():
                self.assertEqual(rows["msvc-pgo"]["job_increase"]["percent"], 50)
                self.assertAlmostEqual(rows["msvc-pgo"]["ironrdp_saving"]["percent"], 20)
                self.assertEqual(rows["msvc-pgo"]["break_even_replays"], 5)
                self.assertEqual(rows["llvm-pgo-thin"]["break_even_replays"], 7)
            path = root / "comparison.docx"
            word_document(result, path)
            from docx import Document
            document = Document(path)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(len(document.tables[0].rows), 7)
            text = "\n".join(p.text for p in document.paragraphs)
            self.assertIn("not today's shipped x64 PGO recipe", text)
            self.assertIn("release run release", text)

    def test_rejects_incomplete_pilot_and_wrong_archive(self):
        with tempfile.TemporaryDirectory() as tmp:
            args = fixture(Path(tmp))
            jobs = args[2]
            jobs["jobs"][0]["conclusion"] = "failure"
            with self.assertRaisesRegex(ValueError, "unfinished, failed"):
                analyze(*args)
            jobs["jobs"][0]["conclusion"] = "success"
            path = args[1] / "x64-1/index.json"
            original = path.read_text()
            index = json.loads(original)
            index["pilot"] = True
            path.write_text(json.dumps(index))
            with self.assertRaisesRegex(ValueError, "incomplete, pilot"):
                analyze(*args)
            index = json.loads(original)
            index["artifact_evidence"]["baseline"]["archives"]["rustc-nightly-x86_64-pc-windows-msvc.tar.xz"] = "wrong"
            path.write_text(json.dumps(index))
            with self.assertRaisesRegex(ValueError, "archive provenance"):
                analyze(*args)
            index = json.loads(original)
            index["runs"].pop()
            path.write_text(json.dumps(index))
            with self.assertRaisesRegex(ValueError, "incomplete paired"):
                analyze(*args)

    def test_paired_fixed_effect_and_input_checks(self):
        value = change([10, 20, 30], [8, 16, 24], paired=True, reduction=True)
        self.assertAlmostEqual(value["percent"], 20)
        for bound in value["ci95"]:
            self.assertAlmostEqual(bound, 20)
        with self.assertRaises(ValueError):
            change([0], [1])


if __name__ == "__main__":
    unittest.main()
