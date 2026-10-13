"""Report separate, fully qualified self-hosted pilots without promoting them to hosted samples."""

import argparse
from datetime import datetime
import json
import math
from pathlib import Path

from analyze_matched_release import HOSTS, validate_native_block
from build_arm_compilers import digest, save
from matched_release import VARIANTS
from matched_release_probe import one_file, release_runs
from msvc_experiment import equivalent_tools
from native_ci_probe import paired_order
from recover_matched_build import RUNNERS


def analyze(releases, index_path, jobs, protocol):
    index = json.loads(index_path.read_text())
    runner = index["runner"]
    machine = runner["runner_class"]
    if machine not in RUNNERS or runner["runner_name"] != RUNNERS[machine]:
        raise ValueError("unexpected self-hosted machine")
    cores = 64 if machine == "self-hosted-arm64-highcore" else 4
    arch = "x64" if machine == "self-hosted-x64" else "arm64"
    host = HOSTS[arch]
    expected_protocol = {**protocol, "cores": cores, "resource_class": machine}
    if (index.get("selfhost_schema") != 1 or not index["complete"] or not index["pilot"]
            or index["benchmark_attempt"] != "1" or index["protocol"] != expected_protocol
            or index["host"] != host or index["vm"] != 1
            or runner["build_affinity_count"] != cores or runner["build_jobs"] != cores
            or runner["logical_cpus"] < cores):
        raise ValueError("incomplete or incompatible self-hosted pilot")
    runs = release_runs(index["release_runs"], runner)
    cells = {}
    for variant in VARIANTS:
        folder = releases / f"release-{arch}-{variant}-1-evidence"
        path = one_file(folder, "result.json")
        item = json.loads(path.read_text())
        if (not item["complete"] or not item["full_distribution"] or not item["pilot"]
                or item["run_id"] != runs[variant] or item["variant"] != variant
                or item["host"] != host or item["vm"] != 1 or item["attempt"] != "1"
                or any(item.get(key) != value for key, value in runner.items())
                or any(item[key] != protocol[key] for key in ("rust_sha", "llvm_sha", "perf_sha"))):
            raise ValueError("release qualification, source or machine mismatch")
        if not equivalent_tools(item["tools"], index["tools"]):
            raise ValueError("release and benchmark tools differ")
        evidence = index["artifact_evidence"][variant]
        if (index["compiler_metadata"][variant] != item
                or evidence["result_sha256"] != digest(path)):
            raise ValueError("actual compiler metadata provenance mismatch")
        manifest = json.loads(one_file(folder, "dist-manifest.json").read_text())
        hashes = {row["name"]: row["sha256"] for row in manifest}
        required = {f"rustc-nightly-{host}.tar.xz"}
        if variant == "baseline":
            required |= {f"cargo-nightly-{host}.tar.xz", f"rust-std-nightly-{host}.tar.xz"}
        if (set(evidence["archives"]) != required
                or any(hashes.get(name) != value for name, value in evidence["archives"].items())):
            raise ValueError("actual distribution archive provenance mismatch")
        matching = [job for job in jobs["jobs"] if str(job["run_id"]) == runs[variant]
                    and job["name"] == f"{arch} / {variant} / 1"]
        if len(matching) != 1:
            raise ValueError("missing or duplicate selected release job")
        job = matching[0]
        if job["status"] != "completed" or job["conclusion"] != "success" or job["run_attempt"] != 1:
            raise ValueError("unfinished or unsuccessful selected release job")
        start, end = (datetime.fromisoformat(job[key].replace("Z", "+00:00"))
                      for key in ("started_at", "completed_at"))
        total = (end - start).total_seconds()
        production = json.loads(one_file(folder, "distribution-timing.json").read_text())["wall_seconds"]
        qualification = json.loads(one_file(folder, "extracted-dist-tests.json").read_text())
        if not math.isfinite(production) or not 0 < production <= total or qualification["exit_code"]:
            raise ValueError("invalid release timing or qualification")
        maintenance = releases / f"release-{arch}-{variant}-1-maintenance"
        seconds = 0
        if maintenance.exists() or variant != "baseline":
            record = json.loads(one_file(maintenance, "space-recovery.json").read_text())
            recorded = record["finished_unix"] - record["started_unix"]
            steps = [step for step in job["steps"]
                     if step["name"] == "Reclaim qualified self-hosted build space after uploads"]
            if len(steps) != 1 or steps[0]["conclusion"] != "success":
                raise ValueError("missing successful maintenance step")
            first, last = (datetime.fromisoformat(steps[0][key].replace("Z", "+00:00"))
                           for key in ("started_at", "completed_at"))
            seconds = (last - first).total_seconds()
            if (not record["complete"] or not 0 <= recorded <= seconds + 2 or not 0 <= seconds < total
                    or record["result_sha256"] != digest(path)):
                raise ValueError("invalid post-build maintenance accounting")
        cells[variant] = {"production_seconds": production, "job_seconds": total,
                          "maintenance_seconds": seconds, "job_id": job["id"],
                          "release_run": runs[variant]}
    if len(index["runs"]) != len(VARIANTS):
        raise ValueError("expected one complete paired pilot round")
    seen = set()
    for run in index["runs"]:
        variant = run["variant"]
        if (variant not in VARIANTS or variant in seen or run["round"] != 0 or run["warmup"]
                or run["position"] != paired_order(list(VARIANTS), 1, 0).index(variant)):
            raise ValueError("missing or incorrectly ordered pilot block")
        seen.add(variant)
        block = json.loads((index_path.parent / run["path"]).read_text())
        seconds = validate_native_block(block, expected_protocol, variant, 0)
        if not math.isfinite(seconds):
            raise ValueError("non-finite benchmark duration")
        cells[variant]["ironrdp_seconds"] = seconds
    for variant, row in cells.items():
        baseline = cells["baseline"]
        row["production_increase_percent"] = 100 * (row["production_seconds"] / baseline["production_seconds"] - 1)
        row["job_increase_percent"] = 100 * (row["job_seconds"] / baseline["job_seconds"] - 1)
        row["ironrdp_saving_percent"] = 100 * (1 - row["ironrdp_seconds"] / baseline["ironrdp_seconds"])
    return {"machine": machine, "runner": runner, "pilot": True, "repetitions": 1,
            "protocol": expected_protocol, "benchmark_run": index["benchmark_run"], "variants": cells}


def word_document(results, path):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt

    if not 1 <= len(results) <= 2 or len({r["machine"] for r in results}) != len(results):
        raise ValueError("one-page pilot report requires one or two distinct machines")
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.7), Inches(8.3)
    section.top_margin = section.bottom_margin = Inches(0.45)
    section.left_margin = section.right_margin = Inches(0.5)
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10)
    style.paragraph_format.space_after = Pt(5)
    document.add_heading("Rust release cost versus IronRDP savings", 0)
    document.add_paragraph(
        "PRELIMINARY: separate self-hosted pilots, one qualified build per recipe and one paired "
        "IronRDP round. Observed percentages only; no confidence intervals or hosted-runner claims.")
    table = document.add_table(rows=1, cols=6)
    table.style = "Light Shading Accent 1"
    table.autofit = False
    for cell, text in zip(table.rows[0].cells, (
            "Machine / recipe", "Production min\n(increase)", "Whole job min\n(increase)",
            "Cleanup min\nwithin job", "IronRDP sec", "Compile saving")):
        cell.text = text
    labels = {"baseline": "No PGO", "msvc-pgo": "PGO / Microsoft retained",
              "llvm-pgo-thin": "PGO + ThinLTO / LLVM"}
    for result in results:
        machine = f"{result['runner']['runner_name']} / {result['runner']['build_jobs']} CPUs"
        for variant, row in result["variants"].items():
            reference = variant == "baseline"
            values = [
                f"{machine}\n{labels[variant]}",
                f"{row['production_seconds'] / 60:.1f}\n" +
                ("Reference" if reference else f"{row['production_increase_percent']:+.1f}%"),
                f"{row['job_seconds'] / 60:.1f}\n" +
                ("Reference" if reference else f"{row['job_increase_percent']:+.1f}%"),
                f"{row['maintenance_seconds'] / 60:.1f}", f"{row['ironrdp_seconds']:.1f}",
                "Reference" if reference else f"{row['ironrdp_saving_percent']:.1f}%",
            ]
            for cell, text in zip(table.add_row().cells, values):
                cell.text = text
    for row in table.rows:
        for cell, width in zip(row.cells, (3.2, 1.6, 1.6, 1.1, 1.6, 1.6)):
            cell.width = Inches(width)
    document.add_paragraph(
        "Production includes fresh full-tool PGO training and the full distribution, tools/docs/MSI "
        "(also Arm64EC on Arm64), not just compilation. Whole jobs additionally include setup, "
        "verification, uploads and the displayed cleanup. Historical baseline recovery was outside "
        "its job; therefore whole-job increases also include this maintenance-policy difference.")
    document.add_paragraph(
        "IronRDP sums the five original native offline Cargo commands; setup, downloads and correctness "
        "tests are outside timing. All recipes use the actual released compiler, common baseline Cargo/std "
        "and identical native tools within each machine. Positive saving means faster compilation.")
    document.add_paragraph(
        "Limits: controlled no-PGO baseline, not the shipped x64 PGO recipe. No pooling across machines "
        "or comparison of CPU architectures. Pilots do not fulfill the preregistered repeated hosted "
        "study. Timed-out or offline-runner cells are excluded, never treated as completed builds.")
    for result in results:
        ids = ", ".join(f"{v}: {r['release_run']}" for v, r in result["variants"].items())
        document.add_paragraph(
            f"{result['machine']}: {ids}; IronRDP {result['benchmark_run']}. "
            f"Rust {result['protocol']['rust_sha'][:12]}; github.com/marcpems/IronRDP-ci-benchmark")
    document.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--releases", type=Path, required=True,
                        help="Root containing one release-artifact subdirectory per runner class")
    parser.add_argument("--index", type=Path, action="append", required=True)
    parser.add_argument("--jobs", type=Path, required=True)
    parser.add_argument("--protocol", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    jobs, protocol = json.loads(args.jobs.read_text()), json.loads(args.protocol.read_text())
    results = [analyze(args.releases / json.loads(path.read_text())["runner"]["runner_class"],
                       path, jobs, protocol) for path in args.index]
    args.output.mkdir(parents=True, exist_ok=True)
    save(args.output / "selfhost-pilot-results.json", results)
    word_document(results, args.output / "Self-Hosted-Rust-Release-Pilots.docx")


if __name__ == "__main__":
    main()
