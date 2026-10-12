"""Reject incomplete matched studies, then produce the one-page Word comparison."""

import argparse
from datetime import datetime
import json
import math
from pathlib import Path
import random
import statistics

from build_arm_compilers import digest, save
from matched_release import VARIANTS
from matched_release_probe import one_file
from msvc_experiment import equivalent_tools
from native_ci_probe import NATIVE_COMMANDS, paired_order

HOSTS = {"x64": "x86_64-pc-windows-msvc", "arm64": "aarch64-pc-windows-msvc"}


def change(control, treatment, paired=False, reduction=False):
    if len(control) != len(treatment) or not control or min(control + treatment) <= 0:
        raise ValueError("expected positive complete comparison samples")
    sign = -1 if reduction else 1
    estimate = sign * 100 * (statistics.mean(treatment) / statistics.mean(control) - 1)
    rng = random.Random(20261011)
    draws = []
    for _ in range(10000):
        indices = rng.choices(range(len(control)), k=len(control))
        other = indices if paired else rng.choices(range(len(treatment)), k=len(treatment))
        draws.append(sign * 100 * (
            sum(treatment[i] for i in other) / sum(control[i] for i in indices) - 1))
    draws.sort()
    return {"percent": estimate, "ci95": [draws[249], draws[9749]]}


def analyze(releases, benchmarks, jobs, protocol):
    n = protocol["release_repetitions_per_cell"]
    cells, runs, tools = {}, set(), {}
    job_rows = jobs["jobs"]
    expected_names = {f"{arch} / {variant} / {vm}"
                      for arch in HOSTS for variant in VARIANTS for vm in range(1, n + 1)}
    if len(job_rows) != len(expected_names) or {j["name"] for j in job_rows} != expected_names:
        raise ValueError("missing, duplicate, or unexpected release jobs")
    job_by_name = {job["name"]: job for job in job_rows}
    for arch, host in HOSTS.items():
        for vm in range(1, n + 1):
            for variant in VARIANTS:
                folder = releases / f"release-{arch}-{variant}-{vm}-evidence"
                path = one_file(folder, "result.json")
                item = json.loads(path.read_text())
                if item.get("runner_class", "github-hosted") != "github-hosted":
                    raise ValueError("self-hosted results cannot enter the hosted-runner comparison")
                if (not item["complete"] or item["pilot"] or item["host"] != host
                        or item["variant"] != variant or item["vm"] != vm or item["attempt"] != "1"):
                    raise ValueError("invalid, pilot or mismatched release evidence")
                for key in ("rust_sha", "llvm_sha", "perf_sha"):
                    if item[key] != protocol[key]:
                        raise ValueError("release source identity differs from protocol")
                runs.add(item["run_id"])
                if arch in tools and not equivalent_tools(tools[arch], item["tools"]):
                    raise ValueError("native tools changed between release jobs")
                tools[arch] = item["tools"]
                job = job_by_name[f"{arch} / {variant} / {vm}"]
                if (job["status"] != "completed" or job["conclusion"] != "success"
                        or str(job["run_id"]) != item["run_id"] or job["run_attempt"] != 1):
                    raise ValueError("unfinished, failed, or mixed-attempt release job")
                start, end = (datetime.fromisoformat(job[key].replace("Z", "+00:00"))
                              for key in ("started_at", "completed_at"))
                total = (end - start).total_seconds()
                production = json.loads(one_file(folder, "distribution-timing.json").read_text())
                qualification = json.loads(one_file(folder, "extracted-dist-tests.json").read_text())
                if not 0 < production["wall_seconds"] <= total or qualification["exit_code"] != 0:
                    raise ValueError("invalid production timer or failed release qualification")
                manifest_path = one_file(folder, "dist-manifest.json")
                manifest = json.loads(manifest_path.read_text())
                cells[arch, vm, variant] = {
                    "metadata": item, "metadata_sha256": digest(path),
                    "manifest": {r["name"]: r["sha256"] for r in manifest},
                    "job_seconds": total, "production_seconds": production["wall_seconds"],
                    "qualification_seconds": qualification["wall_seconds"],
                    "job_id": job["id"],
                }
    if len(runs) != 1:
        raise ValueError("mixed release workflow runs")
    release_run = runs.pop()
    samples, seen, benchmark_runs = {}, set(), set()
    for path in benchmarks.rglob("index.json"):
        index = json.loads(path.read_text())
        host, vm = index["host"], index["vm"]
        matches = [arch for arch, value in HOSTS.items() if value == host]
        if len(matches) != 1 or vm not in range(1, n + 1):
            raise ValueError("unexpected benchmark cell")
        arch = matches[0]
        if (arch, vm) in seen:
            raise ValueError("duplicate benchmark VM")
        seen.add((arch, vm))
        if (not index["complete"] or index["pilot"] or index["protocol"] != protocol
                or index["release_run"] != release_run or index["benchmark_attempt"] != "1"):
            raise ValueError("incomplete, pilot, or mismatched benchmark")
        benchmark_runs.add(index["benchmark_run"])
        if not equivalent_tools(tools[arch], index["tools"]):
            raise ValueError("benchmark native tools differ from release builders")
        for variant in VARIANTS:
            release = cells[arch, vm, variant]
            if index["compiler_metadata"][variant] != release["metadata"]:
                raise ValueError("benchmark did not use the recorded release compiler")
            evidence = index["artifact_evidence"][variant]
            if evidence["result_sha256"] != release["metadata_sha256"]:
                raise ValueError("release metadata hash mismatch")
            required = {f"rustc-nightly-{host}.tar.xz"}
            if variant == "baseline":
                required |= {f"cargo-nightly-{host}.tar.xz", f"rust-std-nightly-{host}.tar.xz"}
            if set(evidence["archives"]) != required or any(
                checksum != release["manifest"][name] for name, checksum in evidence["archives"].items()
            ):
                raise ValueError("benchmark archive provenance mismatch")
        total_rounds = protocol["warmup_rounds"] + protocol["measured_rounds"]
        expected = {(variant, r) for variant in VARIANTS for r in range(total_rounds)}
        blocks, values = set(), {variant: [] for variant in VARIANTS}
        for run in index["runs"]:
            variant, r = run["variant"], run["round"]
            key = variant, r
            if key not in expected or key in blocks:
                raise ValueError("duplicate or unexpected benchmark block")
            blocks.add(key)
            if (run["warmup"] != (r < protocol["warmup_rounds"])
                    or run["position"] != paired_order(list(VARIANTS), vm, r).index(variant)):
                raise ValueError("incorrect warmup or variant order")
            block = json.loads((path.parent / run["path"]).read_text())
            metadata = block["metadata"]
            if metadata["source_sha"] != protocol["source_sha"] or metadata["variant"] != variant or metadata["round"] != r:
                raise ValueError("benchmark block identity mismatch")
            env = metadata["profile_environment"]
            if env["CARGO_NET_OFFLINE"] != "true" or env["CARGO_BUILD_JOBS"] != "4" or "RUST_TEST_THREADS" in env:
                raise ValueError("timed compilation environment changed")
            rows = block["measurements"]
            if [row["name"] for row in rows] != [name for name, _ in NATIVE_COMMANDS]:
                raise ValueError("missing or reordered native command")
            if any(row["exit_code"] or row["wall_seconds"] <= 0 or row["cpu_seconds"] <= 0
                   or row["affinity"] != [0, 1, 2, 3] or "--frozen" not in row["command"]
                   or row["compiled_artifacts"] <= 0 for row in rows):
                raise ValueError("invalid or failed offline compilation")
            if r == 0:
                tests = block["correctness"]
                if ([row["name"] for row in tests] != [name for name, _ in NATIVE_COMMANDS[1:]]
                        or any(row["exit_code"] or row["environment_overrides"] != {"RUST_TEST_THREADS": "1"}
                               for row in tests)):
                    raise ValueError("missing, failed, or changed correctness commands")
            if not run["warmup"]:
                values[variant].append(sum(row["wall_seconds"] for row in rows))
        if blocks != expected:
            raise ValueError("incomplete paired benchmark rounds")
        for variant in VARIANTS:
            samples[arch, vm, variant] = statistics.mean(values[variant])
    if seen != {(arch, vm) for arch in HOSTS for vm in range(1, n + 1)} or len(benchmark_runs) != 1:
        raise ValueError("incomplete benchmark matrix or mixed benchmark runs")
    result = {"protocol": protocol, "release_run": release_run,
              "benchmark_run": benchmark_runs.pop(), "architectures": {}}
    for arch in HOSTS:
        records = {}
        baseline_jobs = [cells[arch, vm, "baseline"]["job_seconds"] for vm in range(1, n + 1)]
        baseline_native = [samples[arch, vm, "baseline"] for vm in range(1, n + 1)]
        for variant in VARIANTS:
            job_times = [cells[arch, vm, variant]["job_seconds"] for vm in range(1, n + 1)]
            native = [samples[arch, vm, variant] for vm in range(1, n + 1)]
            job_mean, native_mean = statistics.mean(job_times), statistics.mean(native)
            extra, saved = job_mean - statistics.mean(baseline_jobs), statistics.mean(baseline_native) - native_mean
            records[variant] = {
                "job_seconds": job_times, "ironrdp_vm_means": native,
                "mean_job_minutes": job_mean / 60, "mean_ironrdp_seconds": native_mean,
                "mean_production_minutes": statistics.mean(
                    cells[arch, vm, variant]["production_seconds"] for vm in range(1, n + 1)) / 60,
                "job_increase": {"percent": 0, "ci95": [0, 0]} if variant == "baseline"
                else change(baseline_jobs, job_times),
                "ironrdp_saving": change(baseline_native, native, paired=True, reduction=True),
                "break_even_replays": math.ceil(extra / saved) if extra > 0 and saved > 0 else None,
            }
        result["architectures"][arch] = records
    return result


def word_document(result, path):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.7), Inches(8.3)
    section.top_margin = section.bottom_margin = Inches(0.45)
    section.left_margin = section.right_margin = Inches(0.5)
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(10)
    style.paragraph_format.space_after = Pt(5)
    document.add_heading("Rust release cost versus IronRDP compilation savings", 0)
    document.add_paragraph(
        "Matched source, cold full distributions, native standard GitHub Windows runners. "
        "Three independent release builds per recipe/architecture; three measured IronRDP rounds "
        "per matched compiler set after one excluded warmup.")
    labels = {"baseline": "No-PGO baseline", "msvc-pgo": "PGO / MSVC retained",
              "llvm-pgo-thin": "PGO + ThinLTO / LLVM tools"}
    table = document.add_table(rows=1, cols=7)
    table.style = "Light Shading Accent 1"
    headers = ["Architecture / recipe", "Full job min\n(production)", "Job increase\n95% interval",
               "IronRDP\nseconds", "Compile saving\n95% interval", "Break-even\nreplays", "Replicas"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for arch, variants in result["architectures"].items():
        for variant, row in variants.items():
            cells = table.add_row().cells
            extra, saving = row["job_increase"], row["ironrdp_saving"]
            entries = [
                f"{arch} / {labels[variant]}",
                f"{row['mean_job_minutes']:.1f}\n({row['mean_production_minutes']:.1f})",
                "Reference" if variant == "baseline" else
                f"{extra['percent']:+.1f}%\n[{extra['ci95'][0]:+.1f}, {extra['ci95'][1]:+.1f}]",
                f"{row['mean_ironrdp_seconds']:.1f}",
                "Reference" if variant == "baseline" else
                f"{saving['percent']:.1f}%\n[{saving['ci95'][0]:.1f}, {saving['ci95'][1]:.1f}]",
                str(row["break_even_replays"]) if row["break_even_replays"] is not None else "-",
                str(len(row["job_seconds"])),
            ]
            for cell, value in zip(cells, entries):
                cell.text = value
    document.add_paragraph(
        "Interpretation: positive job increase is additional Rust production cost; positive compile "
        "saving is faster IronRDP. Break-even divides added release-job seconds by saved native "
        "IronRDP seconds per replay (point estimates, same runner class; not a billing forecast).")
    document.add_paragraph(
        "Boundaries: full job includes setup, fresh profile training, full tools/docs/MSI, common "
        "extracted-distribution tests and artifact uploads. Parentheses show production alone. "
        "IronRDP is the five original native offline Cargo commands, with common baseline Cargo/std; "
        "downloads and correctness runs are excluded. Arm64 retains Arm64EC.")
    document.add_paragraph(
        "Limits: baseline is a controlled no-PGO full distribution, not today's shipped x64 PGO recipe. "
        "Comparisons are within architecture; runner CPU models can differ. Release intervals resample "
        "independent jobs; workload intervals resample matched VM/build clusters. Three replicas give "
        "limited precision. Serial IronRDP correctness does not resolve the earlier parallel-test failure.")
    document.add_paragraph(
        f"Rust {result['protocol']['rust_sha'][:12]}; release run {result['release_run']}; "
        f"IronRDP run {result['benchmark_run']}. Both workflow runs and the full supporting data: "
        "github.com/marcpems/IronRDP-ci-benchmark")
    document.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--releases", type=Path, required=True)
    parser.add_argument("--benchmarks", type=Path, required=True)
    parser.add_argument("--jobs", type=Path, required=True)
    parser.add_argument("--protocol", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = analyze(args.releases, args.benchmarks, json.loads(args.jobs.read_text()),
                     json.loads(args.protocol.read_text()))
    args.output.mkdir(exist_ok=True, parents=True)
    save(args.output / "results.json", result)
    word_document(result, args.output / "Rust-Release-Cost-vs-IronRDP.docx")


if __name__ == "__main__":
    main()
